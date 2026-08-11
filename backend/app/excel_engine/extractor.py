"""Mapping-driven extraction: reads exactly the cells a fund's confirmed
MappingConfig says matter, and nothing else. Statements marked
`present: false` are skipped entirely so a fund missing a statement type
(e.g. no Schedule of Investments for a debt fund) never breaks the run.
"""

import openpyxl


def _to_number(value):
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        cleaned = value.replace(",", "").replace("$", "").strip()
        cleaned = cleaned.replace("(", "-").replace(")", "")
        try:
            return float(cleaned)
        except ValueError:
            return None
    return None


def extract_workbook(path: str, config: dict) -> dict:
    wb = openpyxl.load_workbook(path, data_only=True)
    try:
        result: dict = {}
        for statement_key, statement_cfg in config.get("statements", {}).items():
            if not statement_cfg.get("present"):
                continue
            sheet_name = statement_cfg.get("sheet_name")
            if sheet_name not in wb.sheetnames:
                result[statement_key] = {"error": f"Sheet '{sheet_name}' not found in workbook"}
                continue
            ws = wb[sheet_name]
            stmt_type = statement_cfg.get("type", "line_items")

            if stmt_type == "line_items":
                result[statement_key] = _extract_line_items(ws, statement_cfg)
            elif stmt_type == "table":
                result[statement_key] = _extract_table(ws, statement_cfg)
            elif stmt_type == "notes":
                result[statement_key] = _extract_notes(ws, statement_cfg)
        return result
    finally:
        wb.close()


def _extract_line_items(ws, statement_cfg: dict) -> dict:
    columns = statement_cfg.get("columns", {})
    items = {}
    for line_item in statement_cfg.get("line_items", []):
        row = line_item["row"]
        values = {}
        for col_role, col_letter in columns.items():
            values[col_role] = _to_number(ws[f"{col_letter}{row}"].value)
        items[line_item["key"]] = values
    return items


def _extract_table(ws, statement_cfg: dict) -> dict:
    table_cfg = statement_cfg.get("table", {})
    columns = table_cfg.get("columns", {})
    start_row = table_cfg.get("start_row")
    end_row = table_cfg.get("end_row")
    rows = []
    if start_row and end_row:
        for r in range(start_row, end_row + 1):
            row_data = {}
            has_value = False
            for field, col_letter in columns.items():
                val = ws[f"{col_letter}{r}"].value
                if val not in (None, ""):
                    has_value = True
                numeric = _to_number(val)
                row_data[field] = numeric if numeric is not None else val
            if has_value:
                rows.append(row_data)

    totals = {}
    total_row = table_cfg.get("total_row")
    if total_row:
        for field, col_letter in columns.items():
            totals[field] = _to_number(ws[f"{col_letter}{total_row}"].value)
    else:
        for field in columns:
            numeric_values = [r[field] for r in rows if isinstance(r.get(field), (int, float))]
            if numeric_values:
                totals[field] = sum(numeric_values)

    return {"rows": rows, "totals": totals}


def _extract_notes(ws, statement_cfg: dict) -> dict:
    items = {}
    for note_item in statement_cfg.get("notes_items", []):
        raw = ws[f"{note_item['column']}{note_item['row']}"].value
        value_type = note_item.get("value_type", "text")
        if value_type == "number":
            items[note_item["key"]] = _to_number(raw)
        elif value_type == "bool":
            items[note_item["key"]] = bool(raw) and str(raw).strip().lower() not in ("0", "false", "no", "")
        else:
            items[note_item["key"]] = raw
    return items
