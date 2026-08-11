"""Auto-detection helper used the first time a fund's workbook is uploaded.

It never claims to produce a final mapping — it produces a *suggestion*
(sheet -> canonical statement guesses, header columns, candidate line-item
rows) that a Preparer reviews and confirms/adjusts in the mapping UI. This
keeps the engine generic: nothing here assumes a fixed set of tabs, and any
sheet it can't confidently classify is simply left for the human to map by
hand.
"""

import re
from typing import Optional

import openpyxl
from openpyxl.utils import get_column_letter

from app.excel_engine.mapping_schema import CANONICAL_STATEMENTS

SHEET_NAME_SYNONYMS: dict[str, list[str]] = {
    "balance_sheet": ["bs lead", "balance sheet", "statement of assets", "soalpc", "bs"],
    "income_statement": ["is lead", "income statement", "statement of operations", "is"],
    "schedule_of_investments": ["soi lead", "schedule of investments", "soi"],
    "statement_of_changes_in_partners_capital": [
        "scpc lead",
        "statement of changes",
        "changes in partners capital",
        "scpc",
    ],
    "statement_of_cash_flows": ["socf lead", "cash flow", "statement of cash flows", "socf"],
    "notes": ["notes"],
}

IGNORED_SHEET_PATTERNS = ["contacts and links", "process", "cover", "instructions"]

HEADER_KEYWORDS = {
    "unadjusted": "unadjusted",
    "adjustments": "adjustments",
    "adjusted": "adjusted",
    "prior_year": "prior year",
}


def _normalize(text: str) -> str:
    return re.sub(r"\s+", " ", text.strip().lower())


def list_sheets(path: str) -> list[str]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        return wb.sheetnames
    finally:
        wb.close()


def suggest_statement_key(sheet_name: str) -> Optional[str]:
    normalized = _normalize(sheet_name)
    for statement_key, synonyms in SHEET_NAME_SYNONYMS.items():
        if any(s in normalized for s in synonyms):
            return statement_key
    return None


def is_ignored_sheet(sheet_name: str) -> bool:
    normalized = _normalize(sheet_name)
    return any(p in normalized for p in IGNORED_SHEET_PATTERNS)


def detect_header_columns(ws, max_row: int = 15) -> dict[str, str]:
    """Scan the top of a sheet for Unadjusted/Adjustments/Adjusted/Prior Year
    style column headers. Returns {role: column_letter}."""
    found: dict[str, str] = {}
    for row in ws.iter_rows(min_row=1, max_row=min(max_row, ws.max_row or max_row)):
        for cell in row:
            if not isinstance(cell.value, str):
                continue
            normalized = _normalize(cell.value)
            for role, keyword in HEADER_KEYWORDS.items():
                if role not in found and keyword in normalized:
                    found[role] = get_column_letter(cell.column)
        if len(found) >= 3:
            break
    return found


def detect_label_rows(ws, max_row: int = 300) -> list[dict]:
    """Return candidate {row, label} entries by picking whichever of the
    first few columns holds the most text labels, then listing its non-empty
    string cells. This is only a shortlist for the UI to choose from."""
    scan_cols = list(range(1, 5))  # A-D
    counts = {c: 0 for c in scan_cols}
    limit = min(max_row, ws.max_row or max_row)
    for row in ws.iter_rows(min_row=1, max_row=limit, max_col=4):
        for cell in row:
            if isinstance(cell.value, str) and cell.value.strip():
                counts[cell.column] += 1
    if not any(counts.values()):
        return []
    best_col = max(counts, key=counts.get)

    candidates = []
    for row in ws.iter_rows(min_row=1, max_row=limit, min_col=best_col, max_col=best_col):
        cell = row[0]
        if isinstance(cell.value, str) and cell.value.strip():
            candidates.append({"row": cell.row, "label": cell.value.strip()})
    return candidates


def infer_mapping(path: str) -> dict:
    wb = openpyxl.load_workbook(path, data_only=True)
    try:
        sheet_names = wb.sheetnames
        statements: dict = {}
        warnings: list[str] = []
        claimed_sheets: dict[str, str] = {}

        for sheet_name in sheet_names:
            if is_ignored_sheet(sheet_name):
                continue
            statement_key = suggest_statement_key(sheet_name)
            if not statement_key:
                warnings.append(
                    f"Sheet '{sheet_name}' did not match a known statement type; map it manually if needed."
                )
                continue
            if statement_key in claimed_sheets:
                warnings.append(
                    f"Both '{claimed_sheets[statement_key]}' and '{sheet_name}' look like "
                    f"'{statement_key}'; keeping the first and ignoring the second."
                )
                continue
            claimed_sheets[statement_key] = sheet_name

            ws = wb[sheet_name]
            columns = detect_header_columns(ws)
            label_rows = detect_label_rows(ws)
            statements[statement_key] = {
                "present": True,
                "sheet_name": sheet_name,
                "type": "table" if statement_key == "schedule_of_investments" else "line_items",
                "columns": columns,
                "line_items": [],
                "detected_rows": label_rows[:200],
            }
            if not columns:
                warnings.append(
                    f"Could not auto-detect Unadjusted/Adjustments/Adjusted/Prior Year "
                    f"columns on '{sheet_name}'; set them manually."
                )

        for statement_key in CANONICAL_STATEMENTS:
            if statement_key not in statements:
                statements[statement_key] = {"present": False}
                warnings.append(
                    f"No sheet matched '{statement_key}'. Leaving it marked as not present for "
                    f"this fund — confirm this is expected (e.g. a debt fund with no SOI)."
                )

        return {
            "suggested_config": {"statements": statements, "cross_checks": []},
            "sheet_names": sheet_names,
            "warnings": warnings,
        }
    finally:
        wb.close()
